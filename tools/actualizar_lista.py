#!/usr/bin/env python3
"""
Actualiza datos.json a partir de un nuevo export de "Lista Precios de Ventas"
(Menu Ventas > Lista Precios Venta > Refrescar > Exportar > Exportar como excel).

Uso:
    python3 tools/actualizar_lista.py /ruta/al/nuevo_export.xlsx

Qué hace:
  - Lee el Excel nuevo (encabezado Id/Articulo/Precio Venta U$S/Disponible Global,
    puede variar la fila exacta del encabezado y hay una fila de totales al final
    que se descarta automáticamente).
  - Normaliza espacios/caracteres raros en los nombres de artículo para poder
    cruzarlos de forma confiable con los nombres ya existentes en datos.json
    (los exports de este sistema a veces traen doble espacio o caracteres
    invisibles en el mismo artículo de una semana a otra).
  - Para los artículos que ya existen: actualiza precio y disponible, mantiene
    rubro/compañía/activo/unidad tal cual estaban clasificados.
  - Para artículos nuevos (nombre normalizado no visto antes): los agrega con
    rubro/compañía/activo/unidad en blanco para completar a mano (o pasárselos
    a Claude, que puede inferir varios por el nombre/rubro de productos
    parecidos ya clasificados).
  - Los artículos que ya no aparecen en el export se eliminan de datos.json
    (se listan en la salida para que quede registro de qué se dio de baja).
  - version = fecha-1 (se puede ajustar a mano si se corre más de una vez el
    mismo día), fecha = hoy, lista = se mantiene igual al valor anterior.
  - Imprime un resumen (agregados/removidos/cambios de precio) para pegar en
    el mensaje de reporte.

No hace commit ni push — eso se hace aparte con git, después de revisar el
resumen.
"""
import json
import re
import sys
import datetime
import openpyxl

REPO = "/home/claude/lista-precios-rycdas"


def norm(s):
    s = s.replace("\xa0", " ").replace("\x81", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 5)]
        if vals[0] == "Id" and vals[1] and "articulo" in str(vals[1]).lower():
            return r
    raise RuntimeError("No encontré la fila de encabezado (Id/Articulo/...) en las primeras 30 filas")


def main(xlsx_path):
    with open(f"{REPO}/datos.json", encoding="utf-8") as f:
        d = json.load(f)

    old_by_norm = {}
    for it in d["items"]:
        old_by_norm[norm(it[0])] = it

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = find_header_row(ws)

    new_rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        idv, art, price, disp = [ws.cell(row=r, column=c).value for c in range(1, 5)]
        if art is None:
            continue  # fila de totales u otra fila sin nombre de artículo
        new_rows.append((norm(art), art, price, disp))

    new_items = []
    seen = set()
    added, changed = [], []
    for key, orig, price, disp in new_rows:
        seen.add(key)
        if key in old_by_norm:
            old = old_by_norm[key]
            new_items.append([old[0], price, disp, old[3], old[4], old[5], old[6]])
            if old[1] != price or old[2] != disp:
                changed.append(old[0])
        else:
            new_items.append([key, price, disp, "", "", "", ""])
            added.append(key)

    removed = [old_by_norm[k][0] for k in old_by_norm if k not in seen]
    new_items.sort(key=lambda x: x[0])

    out = {
        "version": datetime.date.today().isoformat() + "-1",
        "lista": d["lista"],
        "fecha": datetime.date.today().isoformat(),
        "items": new_items,
    }
    with open(f"{REPO}/datos.json", "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

    print(f"Artículos antes: {len(d['items'])}  →  ahora: {len(new_items)}")
    print(f"Agregados ({len(added)}): sin clasificar, completar rubro/compañía/activo/unidad")
    for a in added:
        print("  +", a)
    print(f"Dados de baja ({len(removed)}):")
    for r in removed:
        print("  -", r)
    print(f"Precio/disponible actualizado en {len(changed)} artículos")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("uso: python3 actualizar_lista.py /ruta/al/export.xlsx")
        sys.exit(1)
    main(sys.argv[1])
